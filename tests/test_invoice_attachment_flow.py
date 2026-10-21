import asyncio
from datetime import date
from unittest.mock import AsyncMock, patch

import jwt
import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from pydantic import ValidationError
from postgrest.exceptions import APIError

from app.api.v1.endpoints.reports import (
    extract_bearer_token,
    generate_szamlamelleklet,
)
from app.schemas.reports import SzamlamellekletRequest
from app.services.excel_service import ExcelReportService
from app.services.supabase_service import SupabaseDataService


COMPANY_ID = "10000000-0000-0000-0000-000000000001"
CLIENT_ID = "20000000-0000-0000-0000-000000000001"


def make_report(**overrides):
    report = {
        "company_id": COMPANY_ID,
        "client_id": CLIENT_ID,
        "client_code": "A/B",
        "client_name": "=HYPERLINK(\"https://example.test\", \"Acme\")",
        "invoice_attachment_language": "en",
        "used_hours": 2.5,
        "available_hours_per_month": 40.0,
        "hours_from_previous_month": 2.0,
        "difference_hours": 39.5,
        "entries": [
            {
                "project_name": "Project",
                "task_description": "=WEBSERVICE(\"https://example.test\")",
                "hours": 2.5,
                "work_date": "2026-10-01",
            }
        ],
    }
    report.update(overrides)
    return report


def test_request_rejects_invalid_ranges_duplicates_and_partial_month_updates():
    with pytest.raises(ValidationError):
        SzamlamellekletRequest(
            company_id=COMPANY_ID,
            client_ids=[CLIENT_ID, CLIENT_ID],
            start_date="2026-10-01",
            end_date="2026-10-31",
        )

    with pytest.raises(ValidationError):
        SzamlamellekletRequest(
            company_id=COMPANY_ID,
            client_ids=[CLIENT_ID],
            start_date="2026-10-31",
            end_date="2026-10-01",
        )

    with pytest.raises(ValidationError):
        SzamlamellekletRequest(
            company_id=COMPANY_ID,
            client_ids=[CLIENT_ID],
            start_date="2026-10-02",
            end_date="2026-10-31",
            update_remaining_hours=True,
        )

    request = SzamlamellekletRequest(
        company_id=COMPANY_ID,
        client_ids=[CLIENT_ID],
        start_date="2026-10-01",
        end_date="2026-10-31",
        update_remaining_hours=True,
    )
    assert request.start_date == date(2026, 10, 1)


def test_bearer_header_returns_401_instead_of_validation_error():
    for header in (None, "", "Basic token", "Bearer", "Bearer  ", "Bearer a b"):
        with pytest.raises(HTTPException) as exc_info:
            extract_bearer_token(header)
        assert exc_info.value.status_code == 401
        assert exc_info.value.headers == {"WWW-Authenticate": "Bearer"}

    assert extract_bearer_token("bearer valid-token") == "valid-token"


def test_invoice_attachment_access_requires_an_active_company_admin():
    token = jwt.encode(
        {"sub": "30000000-0000-0000-0000-000000000001"},
        "test-key-with-at-least-thirty-two-bytes",
    )

    class ProfileQuery:
        def __init__(self, profile):
            self.profile = profile

        def __getattr__(self, _name):
            return lambda *_args, **_kwargs: self

        async def execute(self):
            return type("Response", (), {"data": [self.profile]})()

    class ProfileClient:
        def __init__(self, profile):
            self.profile = profile

        def table(self, table_name):
            assert table_name == "profiles"
            return ProfileQuery(self.profile)

    asyncio.run(
        SupabaseDataService._assert_invoice_attachment_access_with_client(
            ProfileClient(
                {
                    "id": "30000000-0000-0000-0000-000000000001",
                    "company_id": COMPANY_ID,
                    "role": "company_admin",
                    "is_active": True,
                }
            ),
            token,
            COMPANY_ID,
        )
    )

    with pytest.raises(HTTPException) as exc_info:
        asyncio.run(
            SupabaseDataService._assert_invoice_attachment_access_with_client(
                ProfileClient(
                    {
                        "id": "30000000-0000-0000-0000-000000000001",
                        "company_id": COMPANY_ID,
                        "role": "employee",
                        "is_active": True,
                    }
                ),
                token,
                COMPANY_ID,
            )
        )
    assert exc_info.value.status_code == 403


def test_workbook_localizes_each_sheet_and_neutralizes_formulas():
    reports = [
        make_report(),
        make_report(
            client_id="20000000-0000-0000-0000-000000000002",
            client_code="AB",
            client_name="Hungarian Client",
            invoice_attachment_language="hu",
            entries=[
                {
                    "project_name": "Projekt",
                    "task_description": "+DANGEROUS",
                    "hours": 1,
                    "work_date": "2026-10-02",
                }
            ],
        ),
    ]

    workbook_stream = asyncio.run(
        ExcelReportService().generate_szamlamelleklet(
            reports,
            start_date=date(2026, 10, 1),
            end_date=date(2026, 10, 31),
        )
    )
    workbook = load_workbook(workbook_stream, data_only=False)

    assert workbook.sheetnames == ["AB", "AB (2)"]
    assert workbook["AB"]["C11"].value == "October 2026"
    assert workbook["AB (2)"]["C11"].value == "2026. október"
    assert workbook["AB"]["C6"].value.startswith("'=")
    assert workbook["AB"]["C17"].value.startswith("'=")
    assert workbook["AB (2)"]["C17"].value.startswith("'+")
    assert workbook["AB"].sheet_view.showGridLines is False
    assert workbook["AB"].page_setup.orientation == "landscape"
    assert workbook["AB"].page_setup.fitToWidth == 1
    assert workbook["AB"].print_title_rows == "$16:$16"
    assert workbook.calculation.fullCalcOnLoad is True


def test_hours_use_accounting_rounding_and_match_the_exported_entry_sum():
    assert SupabaseDataService._round_decimal(
        SupabaseDataService._decimal_or_zero("2.345")
    ) == 2.35


def test_company_logo_is_downloaded_once_for_multi_sheet_workbook():
    svg_logo = (
        b'<svg xmlns="http://www.w3.org/2000/svg" width="2" height="2">'
        b'<rect width="2" height="2" fill="white"/></svg>'
    )
    fetch_logo = AsyncMock(return_value=svg_logo)

    reports = [
        make_report(client_name="One"),
        make_report(
            client_id="20000000-0000-0000-0000-000000000002",
            client_code="TWO",
            client_name="Two",
        ),
    ]
    with patch.object(SupabaseDataService, "fetch_company_logo_bytes", fetch_logo):
        workbook_stream = asyncio.run(
            ExcelReportService().generate_szamlamelleklet(
                reports,
                start_date=date(2026, 10, 1),
                end_date=date(2026, 10, 31),
                user_jwt="fake-jwt",
            )
        )

    assert workbook_stream.getbuffer().nbytes > 0
    fetch_logo.assert_awaited_once_with("fake-jwt", COMPANY_ID)
    workbook = load_workbook(workbook_stream)
    assert all(len(sheet._images) == 1 for sheet in workbook.worksheets)


class FakeResponse:
    def __init__(self, data):
        self.data = data


class SnapshotQuery:
    def __init__(self, rows):
        self.rows = rows

    def __getattr__(self, _name):
        return lambda *_args, **_kwargs: self

    async def execute(self):
        return FakeResponse(self.rows)


class SnapshotClient:
    def __init__(self, rows):
        self.query = SnapshotQuery(rows)

    def table(self, table_name):
        assert table_name == "invoice_attachment_periods"
        return self.query


def test_finalized_snapshot_makes_retry_stable_and_detects_changed_hours():
    snapshot = {
        "client_id": CLIENT_ID,
        "starting_hours": "2.00",
        "available_hours": "40.00",
        "used_hours": "2.50",
        "remaining_hours": "39.50",
    }
    report = make_report(hours_from_previous_month=39.5, difference_hours=77.0)

    with patch.object(
        SupabaseDataService,
        "get_authenticated_client",
        return_value=SnapshotClient([snapshot]),
    ):
        asyncio.run(
            SupabaseDataService.reconcile_finalized_periods(
                "fake-jwt",
                COMPANY_ID,
                "2026-10-01",
                "2026-10-31",
                [report],
            )
        )

    assert report["hours_from_previous_month"] == 2.0
    assert report["difference_hours"] == 39.5

    changed_report = make_report(used_hours=3.0)
    with (
        patch.object(
            SupabaseDataService,
            "get_authenticated_client",
            return_value=SnapshotClient([snapshot]),
        ),
        pytest.raises(HTTPException) as exc_info,
    ):
        asyncio.run(
            SupabaseDataService.reconcile_finalized_periods(
                "fake-jwt",
                COMPANY_ID,
                "2026-10-01",
                "2026-10-31",
                [changed_report],
            )
        )
    assert exc_info.value.status_code == 409


def test_finalization_conflicts_are_returned_as_http_409():
    class FailingRpc:
        async def execute(self):
            raise APIError(
                {
                    "code": "40001",
                    "message": "Invoice attachment data changed before finalization",
                }
            )

    class FailingRpcClient:
        def rpc(self, function_name, params):
            assert function_name == "finalize_invoice_attachment_period"
            assert params["p_company_id"] == COMPANY_ID
            return FailingRpc()

    with (
        patch.object(
            SupabaseDataService,
            "get_authenticated_client",
            return_value=FailingRpcClient(),
        ),
        pytest.raises(HTTPException) as exc_info,
    ):
        asyncio.run(
            SupabaseDataService.update_remaining_hours_from_reports(
                "fake-jwt",
                COMPANY_ID,
                "2026-10-01",
                "2026-10-31",
                [make_report()],
            )
        )

    assert exc_info.value.status_code == 409


def test_report_endpoint_returns_a_no_store_xlsx_and_runs_finalization():
    fetch_reports = AsyncMock(return_value=[make_report(client_name="Acme")])
    reconcile = AsyncMock(return_value=None)
    finalize = AsyncMock(return_value=None)
    fetch_logo = AsyncMock(return_value=None)

    with (
        patch.object(
            SupabaseDataService,
            "fetch_aggregated_timesheets",
            fetch_reports,
        ),
        patch.object(
            SupabaseDataService,
            "reconcile_finalized_periods",
            reconcile,
        ),
        patch.object(
            SupabaseDataService,
            "update_remaining_hours_from_reports",
            finalize,
        ),
        patch.object(
            SupabaseDataService,
            "fetch_company_logo_bytes",
            fetch_logo,
        ),
    ):
        response = asyncio.run(
            generate_szamlamelleklet(
                SzamlamellekletRequest(
                    company_id=COMPANY_ID,
                    client_ids=[CLIENT_ID],
                    start_date="2026-10-01",
                    end_date="2026-10-31",
                    update_remaining_hours=True,
                ),
                token="fake-jwt",
            )
        )
        reconcile.assert_awaited_once()
        finalize.assert_awaited_once()
        reconcile.reset_mock()
        finalize.reset_mock()

        read_only_response = asyncio.run(
            generate_szamlamelleklet(
                SzamlamellekletRequest(
                    company_id=COMPANY_ID,
                    client_ids=[CLIENT_ID],
                    start_date="2026-10-01",
                    end_date="2026-10-31",
                    update_remaining_hours=False,
                ),
                token="fake-jwt",
            )
        )

    assert response.body.startswith(b"PK")
    assert response.headers["cache-control"] == "no-store"
    assert response.headers["x-content-type-options"] == "nosniff"
    assert (
        response.headers["content-disposition"]
        == 'attachment; filename="invoice_attachments_2026-10-01_2026-10-31.xlsx"'
    )
    assert read_only_response.body.startswith(b"PK")
    reconcile.assert_awaited_once()
    finalize.assert_not_awaited()
