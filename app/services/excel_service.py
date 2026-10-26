import io
import logging
import math
import os
import re
from calendar import monthrange
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, List, Optional
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image as PILImage

from app.services.supabase_service import SupabaseDataService

logger = logging.getLogger(__name__)


class ExcelReportService:
    DEFAULT_LOGO_PATH = Path(__file__).resolve().parents[2] / "assets" / "ecovis_logo.png"
    INVALID_XML_CHARACTERS_RE = re.compile(
        r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]"
    )
    MAX_LOGO_PIXELS = 20_000_000
    MAX_NORMALIZED_LOGO_SIZE = (1448, 294)
    TEXTS = {
        "hu": {
            "title": "Számlamelléklet (teljesítési igazolás)",
            "statement": "Szerződésünk 4. pontja szerint csatoljuk az adott elszámolási időszakban igénybe vett tanácsadási szolgáltatásokról szóló kimutatást.",
            "used_hours_sum": "Felhasznált tanácsadói órák",
            "available_hours": "Szerződés szerint rendelkezésre álló óraszám",
            "previous_hours": "Korábbi időszaki órák",
            "difference": "Különbözet",
            "task": "Feladat",
            "hours": "Időráfordítás (óra)",
        },
        "en": {
            "title": "Invoice attachment (certificate of completion)",
            "statement": "According to point 4 of our contract, we are attaching a statement of the consulting services used in the given accounting period.",
            "used_hours_sum": "Used consulting hours",
            "available_hours": "Contracted available hours",
            "previous_hours": "Hours carried over from the previous period",
            "difference": "Difference",
            "task": "Task description",
            "hours": "Time spent (hours)",
        },
    }
    MONTH_NAMES = {
        "hu": (
            "január",
            "február",
            "március",
            "április",
            "május",
            "június",
            "július",
            "augusztus",
            "szeptember",
            "október",
            "november",
            "december",
        ),
        "en": (
            "January",
            "February",
            "March",
            "April",
            "May",
            "June",
            "July",
            "August",
            "September",
            "October",
            "November",
            "December",
        ),
    }

    def __init__(self, logo_path: Optional[str] = None):
        """Initialize the service with an optional path to the Ecovis logo."""

        self.logo_path: Optional[str] = None
        self.set_logo_path(logo_path)

    def set_logo_path(self, logo_path: Optional[str] = None) -> "ExcelReportService":
        """Store a logo path for the next workbook generation."""
        if logo_path is None:
            logo_path = str(self.DEFAULT_LOGO_PATH)

        self.logo_path = logo_path
        return self

    def _resolve_logo_path(self, logo_path: Optional[str] = None) -> Optional[str]:
        """Resolve the logo path lazily."""

        if logo_path is not None and os.path.exists(logo_path):
            return logo_path
        if self.logo_path and os.path.exists(self.logo_path):
            return self.logo_path

        default_path = str(self.DEFAULT_LOGO_PATH)
        if os.path.exists(default_path):
            return default_path
        return None

    @staticmethod
    def _add_logo_to_sheet(
        ws, logo_path: Optional[str], cell_anchor: str = "A1"
    ) -> None:
        """Adds company logo at the designated anchor cell matching reference dimensions."""
        if not logo_path or not os.path.exists(logo_path):
            logger.warning("No logo path available for sheet %s", ws.title)
            return

        try:
            with open(logo_path, "rb") as logo_file:
                ExcelReportService._add_logo_bytes_to_sheet(
                    ws, logo_file.read(), cell_anchor
                )
        except OSError as exc:
            logger.exception(
                "Failed to read logo image for sheet %s from %s: %s",
                ws.title,
                logo_path,
                exc,
            )

    @classmethod
    def _normalize_logo_bytes(cls, logo_bytes: bytes | None) -> bytes | None:
        if not logo_bytes:
            return None

        try:
            normalized_bytes = logo_bytes
            if b"<svg" in logo_bytes[:512].lower():
                try:
                    import cairosvg
                except ImportError:
                    logger.warning(
                        "SVG logo support is unavailable; omitting the company logo"
                    )
                    return None
                normalized_bytes = cairosvg.svg2png(
                    bytestring=logo_bytes,
                    output_width=cls.MAX_NORMALIZED_LOGO_SIZE[0],
                    output_height=cls.MAX_NORMALIZED_LOGO_SIZE[1],
                    unsafe=False,
                )

            with PILImage.open(io.BytesIO(normalized_bytes)) as pil_img:
                if pil_img.width * pil_img.height > cls.MAX_LOGO_PIXELS:
                    raise ValueError("logo dimensions exceed the safe image limit")
                pil_img.verify()
            with PILImage.open(io.BytesIO(normalized_bytes)) as pil_img:
                pil_img = pil_img.convert("RGBA")
                pil_img.thumbnail(cls.MAX_NORMALIZED_LOGO_SIZE)
                image_bytes = io.BytesIO()
                pil_img.save(image_bytes, format="PNG")
                return image_bytes.getvalue()
        except Exception as exc:
            logger.warning("Ignoring invalid company logo: %s", exc)
            return None

    @staticmethod
    def _add_normalized_logo_to_sheet(
        ws, normalized_logo_bytes: bytes | None, cell_anchor: str = "A1"
    ) -> None:
        if not normalized_logo_bytes:
            return

        try:
            image_bytes = io.BytesIO(normalized_logo_bytes)
            img = Image(image_bytes)
            img.width = 724
            img.height = 147
            ws.add_image(img, cell_anchor)
        except Exception as exc:
            logger.warning("Ignoring invalid logo for sheet %s: %s", ws.title, exc)

    @classmethod
    def _add_logo_bytes_to_sheet(
        cls, ws, logo_bytes: bytes | None, cell_anchor: str = "A1"
    ) -> None:
        cls._add_normalized_logo_to_sheet(
            ws,
            cls._normalize_logo_bytes(logo_bytes),
            cell_anchor,
        )

    @staticmethod
    def _safe_cell_text(value: Any) -> str:
        text = ExcelReportService.INVALID_XML_CHARACTERS_RE.sub("", str(value or ""))
        if text.lstrip().startswith(("=", "+", "-", "@")):
            return f"'{text}"
        return text

    @staticmethod
    def _safe_sheet_title(raw_title: Any, existing_titles: list[str]) -> str:
        title = ExcelReportService.INVALID_XML_CHARACTERS_RE.sub("", str(raw_title or ""))
        base = re.sub(r"[\\/*?:\[\]]", "", title).strip()
        base = base[:31] or "Client"
        existing = {title.casefold() for title in existing_titles}
        candidate = base
        suffix_number = 2
        while candidate.casefold() in existing:
            suffix = f" ({suffix_number})"
            candidate = f"{base[: 31 - len(suffix)]}{suffix}"
            suffix_number += 1
        return candidate

    @classmethod
    def _format_period_text(
        cls, start_date: date, end_date: date, language: str
    ) -> str:
        language = language if language in cls.MONTH_NAMES else "hu"
        is_full_month = (
            start_date.year == end_date.year
            and start_date.month == end_date.month
            and start_date.day == 1
            and end_date.day == monthrange(start_date.year, start_date.month)[1]
        )
        if is_full_month:
            month_name = cls.MONTH_NAMES[language][start_date.month - 1]
            return (
                f"{start_date.year}. {month_name}"
                if language == "hu"
                else f"{month_name} {start_date.year}"
            )

        if language == "hu":
            return (
                f"{start_date.year}. {start_date.month:02d}. {start_date.day:02d}. – "
                f"{end_date.year}. {end_date.month:02d}. {end_date.day:02d}."
            )

        start_month = cls.MONTH_NAMES["en"][start_date.month - 1][:3]
        end_month = cls.MONTH_NAMES["en"][end_date.month - 1][:3]
        return (
            f"{start_month} {start_date.day}, {start_date.year} – "
            f"{end_month} {end_date.day}, {end_date.year}"
        )

    async def generate_szamlamelleklet(
        self,
        client_reports: List[Dict[str, Any]],
        period_text: Optional[str] = None,
        issue_date_text: Optional[str] = None,
        start_date: date | None = None,
        end_date: date | None = None,
        user_jwt: Optional[str] = None,
    ) -> io.BytesIO:
        """
        Generates a multi-sheet Excel workbook where each client has a dedicated sheet
        matching the exact layout, granular per-entry listing, formulas, and visual styling
        of szamlamelleklet_szeptember.xlsx.
        """

        wb = Workbook()
        wb.remove(wb.active)  # type: ignore
        wb.calculation.calcMode = "auto"
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True

        # Typography & Color Palettes
        font_client_name = Font(
            name="Calibri", size=12, bold=False, italic=True, color="FFC00000"
        )
        font_doc_title = Font(name="Calibri", size=14, bold=True, italic=False)
        font_period = Font(name="Calibri", size=11, bold=False, italic=False)
        font_issue_date = Font(
            name="Calibri", size=11, bold=False, italic=True, color="FFC00000"
        )
        font_body = Font(name="Calibri", size=11, bold=False, italic=False)
        font_header = Font(name="Calibri", size=11, bold=True, italic=False)
        font_summary_lbl = Font(name="Calibri", size=11, bold=True, italic=False)
        font_summary_val = Font(name="Calibri", size=11, bold=False, italic=False)

        fill_header = PatternFill(
            start_color="FFF2F2F2", end_color="FFF2F2F2", fill_type="solid"
        )

        align_center = Alignment(horizontal="center")
        align_right = Alignment(horizontal="right")
        align_left = Alignment(horizontal="left")
        align_wrap_left = Alignment(horizontal="left", wrap_text=True)
        align_entry_task = Alignment(wrap_text=True)

        gray_border_side = Side(border_style="thin", color="FF999999")
        table_border = Border(
            left=gray_border_side,
            right=gray_border_side,
            top=gray_border_side,
            bottom=gray_border_side,
        )
        client_border = Border(bottom=Side(border_style="medium", color="FF000000"))

        now = datetime.now(ZoneInfo("Europe/Budapest"))
        effective_issue_date = (
            issue_date_text or f"Budapest, {now.strftime('%Y. %m. %d.')}"
        )
        company_logo_cache: dict[str, bytes | None] = {}

        for report in client_reports:
            client_code = report.get("client_code", "CLIENT")
            safe_title = self._safe_sheet_title(client_code, wb.sheetnames)
            client_name = self._safe_cell_text(report.get("client_name", ""))
            entries = report.get("entries", [])
            logo_path = report.get("logo_path")
            logo_bytes: bytes | None = None

            logo_source_id = report.get("company_id")
            if not logo_source_id:
                logger.warning(
                    "No company_id found for sheet %s; cannot fetch logo from Supabase",
                    safe_title,
                )
            if not logo_path and user_jwt and logo_source_id:
                logo_cache_key = str(logo_source_id)
                if logo_cache_key not in company_logo_cache:
                    company_logo_cache[logo_cache_key] = self._normalize_logo_bytes(
                        await SupabaseDataService.fetch_company_logo_bytes(
                            user_jwt,
                            logo_cache_key,
                        )
                    )
                logo_bytes = company_logo_cache[logo_cache_key]

            language = str(report.get("invoice_attachment_language", "hu")).lower()
            text = self.TEXTS.get(language, self.TEXTS["hu"])

            if start_date and end_date:
                current_period = self._format_period_text(
                    start_date, end_date, language
                )
            else:
                current_period = period_text
            if not current_period:
                current_period = (
                    f"{now.year}. {now.month}. hónap"
                    if language == "hu"
                    else f"{now.year}-{now.month}"
                )

            resolved_logo_path = self._resolve_logo_path(logo_path)

            ws = wb.create_sheet(title=safe_title)

            # Column Dimensions (exact match)
            ws.column_dimensions["A"].width = 10.0
            ws.column_dimensions["C"].width = 90.0
            ws.column_dimensions["D"].width = 21.0
            ws.column_dimensions["E"].width = 10.0

            # Row Heights (exact match)
            ws.row_dimensions[1].height = 45.0
            ws.row_dimensions[6].height = 15.5
            ws.row_dimensions[10].height = 18.5
            ws.row_dimensions[13].height = 29.0

            # Logo Placement at A1
            if logo_path and resolved_logo_path:
                self._add_logo_to_sheet(ws, resolved_logo_path, "A1")
            elif logo_bytes:
                self._add_normalized_logo_to_sheet(ws, logo_bytes, "A1")
            elif resolved_logo_path:
                self._add_logo_to_sheet(ws, resolved_logo_path, "A1")

            # Document Metadata Headers (Rows 6-13, Column C)
            cell_c6 = ws.cell(row=6, column=3, value=client_name)
            cell_c6.font = font_client_name
            cell_c6.alignment = Alignment(horizontal="left", vertical="center")
            cell_c6.border = client_border

            cell_c10 = ws.cell(row=10, column=3, value=text["title"])
            cell_c10.font = font_doc_title
            cell_c10.alignment = align_wrap_left

            cell_c11 = ws.cell(
                row=11, column=3, value=self._safe_cell_text(current_period)
            )
            cell_c11.font = font_period
            cell_c11.alignment = align_wrap_left

            cell_c12 = ws.cell(
                row=12,
                column=3,
                value=self._safe_cell_text(effective_issue_date),
            )
            cell_c12.font = font_issue_date
            cell_c12.alignment = align_wrap_left

            cell_c13 = ws.cell(row=13, column=3, value=text["statement"])
            cell_c13.font = font_body
            cell_c13.alignment = align_wrap_left

            # Table Header (Row 16)
            hdr_c16 = ws.cell(row=16, column=3, value=text["task"])
            hdr_d16 = ws.cell(row=16, column=4, value=text["hours"])
            for hdr in (hdr_c16, hdr_d16):
                hdr.font = font_header
                hdr.fill = fill_header
                hdr.alignment = align_center
                hdr.border = table_border

            # Data Rows (Row 17+)
            first_entry_row = 17
            current_row = first_entry_row

            # Iterate over granular logged entries, not grouped projects
            for entry in entries:
                # Prefer entry-level description/task text; fallback to task or project name
                task_description = (
                    entry.get("task_description")
                    or entry.get("description")
                    or entry.get("task_name")
                    or entry.get("task")
                    or entry.get("project_name", "")
                )
                raw_hours = entry.get("hours", 0) or 0
                hours_val = (
                    int(raw_hours)
                    if float(raw_hours).is_integer()
                    else round(float(raw_hours), 2)
                )

                safe_task_description = self._safe_cell_text(task_description)
                cell_task = ws.cell(
                    row=current_row,
                    column=3,
                    value=safe_task_description,
                )
                cell_task.font = font_body
                cell_task.alignment = align_entry_task
                cell_task.border = table_border

                cell_hours = ws.cell(row=current_row, column=4, value=hours_val)
                cell_hours.font = font_body
                cell_hours.alignment = align_right
                cell_hours.number_format = "0.00"
                cell_hours.border = table_border

                estimated_lines = sum(
                    max(1, math.ceil(len(line) / 90))
                    for line in safe_task_description.splitlines() or [""]
                )
                ws.row_dimensions[current_row].height = min(
                    409.0, max(15.0, estimated_lines * 15.0)
                )

                current_row += 1

            last_entry_row = max(first_entry_row, current_row - 1)

            # Summary Rows (Always rendered unconditionally in the reference template)
            total_used_row = current_row
            cell_total_label = ws.cell(
                row=total_used_row, column=3, value=text["used_hours_sum"]
            )
            cell_total_label.font = font_summary_lbl
            cell_total_label.border = table_border

            sum_formula = (
                f"=SUM(D{first_entry_row}:D{last_entry_row})" if entries else 0
            )
            cell_total_val = ws.cell(row=total_used_row, column=4, value=sum_formula)
            cell_total_val.font = font_summary_val
            cell_total_val.number_format = "0.00"
            cell_total_val.border = table_border

            # Contracted Available Hours
            raw_available = report.get("available_hours_per_month")
            available_val = (
                0
                if raw_available is None
                else (
                    int(raw_available)
                    if float(raw_available).is_integer()
                    else float(raw_available)
                )
            )
            avail_row = total_used_row + 1
            cell_avail_lbl = ws.cell(
                row=avail_row, column=3, value=text["available_hours"]
            )
            cell_avail_lbl.font = font_summary_lbl
            cell_avail_lbl.border = table_border

            cell_avail_val = ws.cell(row=avail_row, column=4, value=available_val)
            cell_avail_val.font = font_summary_val
            cell_avail_val.number_format = "0.00"
            cell_avail_val.border = table_border

            # Older Period Hours
            raw_prev = report.get("hours_from_previous_month")
            previous_val = (
                0
                if raw_prev is None
                else (
                    int(raw_prev) if float(raw_prev).is_integer() else float(raw_prev)
                )
            )
            prev_row = avail_row + 1
            cell_prev_lbl = ws.cell(
                row=prev_row, column=3, value=text["previous_hours"]
            )
            cell_prev_lbl.font = font_summary_lbl
            cell_prev_lbl.border = table_border

            cell_prev_val = ws.cell(row=prev_row, column=4, value=previous_val)
            cell_prev_val.font = font_summary_val
            cell_prev_val.number_format = "0.00"
            cell_prev_val.border = table_border

            # Difference Formula Row
            diff_row = prev_row + 1
            cell_diff_lbl = ws.cell(row=diff_row, column=3, value=text["difference"])
            cell_diff_lbl.font = font_summary_lbl
            cell_diff_lbl.border = table_border

            diff_formula = f"=D{avail_row}-D{total_used_row}+D{prev_row}"
            cell_diff_val = ws.cell(row=diff_row, column=4, value=diff_formula)
            cell_diff_val.font = font_summary_val
            cell_diff_val.number_format = "0.00"
            cell_diff_val.border = table_border

            # Keep each attachment printable at one page wide while allowing long
            # entry lists to continue vertically with their table header repeated.
            ws.sheet_view.showGridLines = False
            ws.print_area = f"A1:E{diff_row}"
            ws.print_title_rows = "16:16"
            ws.page_setup.orientation = "landscape"
            ws.page_setup.paperSize = ws.PAPERSIZE_A4
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.print_options.horizontalCentered = True
            ws.page_margins.left = 0.25
            ws.page_margins.right = 0.25
            ws.page_margins.top = 0.5
            ws.page_margins.bottom = 0.5
            ws.page_margins.header = 0.2
            ws.page_margins.footer = 0.2

        file_stream = io.BytesIO()
        wb.save(file_stream)
        file_stream.seek(0)
        return file_stream
